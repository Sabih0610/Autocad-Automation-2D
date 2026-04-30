"""
Phase 3 — Tasks 3.3 + 3.4 + 3.6 + 3.7.

Read a P&ID drawing, find every LINE_BLOCK_TEST, extract attributes,
and write the result to an Excel line list.

Run from the project root:
    python -m src.use_cases.line_list_extract
"""

import sys
import time
from datetime import datetime
from pathlib import Path

import win32com.client

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from src.logging.decorators import log_job


# ---------- config ----------

DRAWINGS_FOLDER = Path(r"E:\RC-Projects")
LINE_BLOCK_NAME = "LINE_BLOCK_TEST"
TARGET_FILE = "pid_001.dwg"

# Column order in the Excel output. Also the set of attributes we read.
LINE_LIST_COLUMNS = ["LINE_NO", "SIZE", "SPEC", "SERVICE", "FROM", "TO"]

# Where to write the Excel output. Sits next to the project source.
OUTPUT_FOLDER = Path(__file__).parent.parent.parent / "outputs"


# ---------- helpers ----------

def get_acad():
    return win32com.client.GetActiveObject("AutoCAD.Application")


def find_blocks_by_name(model_space, block_name):
    for i in range(model_space.Count):
        entity = model_space.Item(i)
        if entity.ObjectName == "AcDbBlockReference" and entity.Name == block_name:
            yield entity


def extract_attributes(block_ref):
    if not block_ref.HasAttributes:
        return {}
    return {
        a.TagString.strip(): a.TextString
        for a in block_ref.GetAttributes()
    }


def normalize_row(attrs, expected):
    return {tag: attrs.get(tag) for tag in expected}


# ---------- excel output ----------

def write_line_list_xlsx(rows, source_dwg_name, output_path):
    """Write rows to a styled Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Line List"

    # Title row
    ws["A1"] = "Line List"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(LINE_LIST_COLUMNS))

    # Metadata row
    ws["A2"] = f"Source: {source_dwg_name}    Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(italic=True, color="555555")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(LINE_LIST_COLUMNS))

    # Header row
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center")
    for col_idx, col_name in enumerate(LINE_LIST_COLUMNS, start=1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Data rows
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, col_name in enumerate(LINE_LIST_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name) or "")

    # Auto-width columns based on content
    for col_idx, col_name in enumerate(LINE_LIST_COLUMNS, start=1):
        max_len = len(col_name)
        for r in rows:
            val = str(r.get(col_name) or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    # Freeze header rows
    ws.freeze_panes = "A5"

    wb.save(output_path)


# ---------- main ----------

@log_job("line_list_extract")
def main():
    target = DRAWINGS_FOLDER / TARGET_FILE
    if not target.exists():
        print(f"ERROR: file not found: {target}")
        sys.exit(1)

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

    try:
        acad = get_acad()
    except Exception as e:
        print(f"ERROR: cannot connect to AutoCAD ({e}).")
        sys.exit(1)

    print(f"Connected to: {acad.Caption}")
    print(f"Opening {target} ...")
    doc = acad.Documents.Open(str(target))
    time.sleep(0.3)

    try:
        blocks = list(find_blocks_by_name(doc.ModelSpace, LINE_BLOCK_NAME))
        print(f"Active drawing: {doc.Name}")
        print(f"Found {len(blocks)} '{LINE_BLOCK_NAME}' block(s)\n")

        if not blocks:
            print("Nothing to extract. Exiting.")
            return

        rows = [normalize_row(extract_attributes(b), LINE_LIST_COLUMNS) for b in blocks]

        # Console preview
        col_widths = {
            tag: max(len(tag), max((len(str(r[tag] or "")) for r in rows), default=0))
            for tag in LINE_LIST_COLUMNS
        }
        header = " | ".join(tag.ljust(col_widths[tag]) for tag in LINE_LIST_COLUMNS)
        print(header)
        print("-" * len(header))
        for r in rows:
            print(" | ".join(str(r[tag] or "").ljust(col_widths[tag]) for tag in LINE_LIST_COLUMNS))

        # Excel output
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        out_filename = f"line_list_{Path(TARGET_FILE).stem}_{timestamp}.xlsx"
        out_path = OUTPUT_FOLDER / out_filename
        write_line_list_xlsx(rows, doc.Name, out_path)

        print(f"\nWrote {len(rows)} row(s) to: {out_path}")

        # Issue report
        missing_report = []
        for idx, r in enumerate(rows, start=1):
            missing = [t for t in LINE_LIST_COLUMNS if not r[t]]
            if missing:
                missing_report.append((idx, missing))

        if missing_report:
            print("\nIssues found:")
            for idx, missing in missing_report:
                print(f"  Block {idx}: missing {missing}")
        else:
            print("All blocks have all expected attributes.")

    finally:
        doc.Close(False)


if __name__ == "__main__":
    main()
