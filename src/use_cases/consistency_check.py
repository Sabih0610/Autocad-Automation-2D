"""
Phase 6 — Consistency Checker.

First version:
- Compare P&ID line blocks against the generated Excel line list.
- No AI yet.
- No isometric yet.
- Read-only against AutoCAD.
- Outputs an Excel mismatch report.

Run from project root:
    python -m src.use_cases.consistency_check

Optional:
    python -m src.use_cases.consistency_check --pid "E:\\RC-Projects\\pid_001.dwg"
    python -m src.use_cases.consistency_check --excel "E:\\RC-Projects\\autocad-ai\\outputs\\line_list_pid_001_....xlsx"

AutoCAD must be running.
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Tuple

import pywintypes
import win32com.client
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from src.logging.decorators import log_job


# ---------- config ----------

RPC_E_CALL_REJECTED = -2147418111

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_FOLDER = PROJECT_ROOT / "outputs"

DRAWINGS_FOLDER = Path(r"E:\RC-Projects")
DEFAULT_PID_DWG = DRAWINGS_FOLDER / "pid_001.dwg"

LINE_BLOCK_NAME = "LINE_BLOCK_TEST"
LINE_COLUMNS = ["LINE_NO", "SIZE", "SPEC", "SERVICE", "FROM", "TO"]
COMPARE_COLUMNS = ["SIZE", "SPEC", "SERVICE", "FROM", "TO"]


# ---------- COM helpers ----------

def is_autocad_busy_error(error: Exception) -> bool:
    """Detect AutoCAD RPC busy/rejected COM errors."""
    if isinstance(error, pywintypes.com_error):
        try:
            return int(error.hresult) == RPC_E_CALL_REJECTED
        except Exception:
            pass

    text = str(error).lower()
    return (
        "call was rejected by callee" in text
        or "rpc_e_call_rejected" in text
        or str(RPC_E_CALL_REJECTED) in text
    )


def is_stale_proxy_error(error: Exception) -> bool:
    """Detect stale AutoCAD COM proxy errors."""
    text = str(error).lower()
    return (
        isinstance(error, AttributeError)
        or "<unknown>" in text
        or "open.name" in text
        or "open.modelspace" in text
    )


def com_retry(
    operation: Callable[[], Any],
    description: str,
    attempts: int = 5,
    delay_seconds: float = 0.5,
):
    """Retry AutoCAD COM operations when AutoCAD is temporarily busy/stale."""
    last_error = None

    for attempt in range(1, attempts + 1):
        try:
            return operation()
        except Exception as e:
            last_error = e

            retryable = is_autocad_busy_error(e) or is_stale_proxy_error(e)

            if not retryable:
                raise

            if attempt >= attempts:
                break

            print(
                f"AutoCAD COM issue during {description}. "
                f"Retrying {attempt}/{attempts - 1}..."
            )
            time.sleep(delay_seconds)

    raise last_error


def get_acad():
    """Connect to running AutoCAD."""
    try:
        return win32com.client.GetActiveObject("AutoCAD.Application")
    except Exception as e:
        print("ERROR: AutoCAD is not running.")
        print("Open AutoCAD, then run this script again.")
        print(f"Details: {type(e).__name__}: {e}")
        sys.exit(1)


def open_drawing_with_retry(acad, drawing_path: Path):
    """
    Open a DWG and return a usable document COM object.

    This retries the whole open process because AutoCAD can return a stale
    Open proxy before doc.Name / doc.ModelSpace are ready.
    """
    last_error = None

    for attempt in range(1, 4):
        doc = None

        try:
            doc = com_retry(
                lambda: acad.Documents.Open(str(drawing_path)),
                f"opening {drawing_path.name}",
                attempts=5,
                delay_seconds=0.7,
            )

            time.sleep(0.7)

            # Touch required properties with retry so we know the proxy is usable.
            _ = com_retry(
                lambda: doc.Name,
                f"reading {drawing_path.name} document name",
                attempts=5,
                delay_seconds=0.5,
            )

            _ = com_retry(
                lambda: doc.ModelSpace,
                f"reading {drawing_path.name} ModelSpace",
                attempts=5,
                delay_seconds=0.5,
            )

            return doc

        except Exception as e:
            last_error = e
            print(
                f"Open attempt {attempt}/3 failed for {drawing_path.name}: "
                f"{type(e).__name__}: {e}"
            )

            if doc is not None:
                try:
                    doc.Close(False)
                except Exception:
                    pass

            time.sleep(1.0)

            # Refresh AutoCAD application proxy before next attempt.
            try:
                acad = win32com.client.GetActiveObject("AutoCAD.Application")
            except Exception:
                pass

    raise last_error


def get_block_name(block_ref) -> str:
    """Get block effective name safely."""
    try:
        return str(block_ref.EffectiveName)
    except Exception:
        try:
            return str(block_ref.Name)
        except Exception:
            return ""


def extract_attributes(block_ref) -> Dict[str, str]:
    """Return {TAG: value} for a block reference."""
    values: Dict[str, str] = {}

    try:
        has_attrs = bool(block_ref.HasAttributes)
    except Exception:
        has_attrs = False

    if not has_attrs:
        return values

    try:
        attributes = com_retry(
            lambda: block_ref.GetAttributes(),
            "reading line block attributes",
            attempts=5,
            delay_seconds=0.4,
        )
    except Exception:
        return values

    for att in attributes:
        try:
            tag = str(att.TagString).strip()
            value = str(att.TextString).strip()
            if tag:
                values[tag] = value
        except Exception:
            continue

    return values


def find_line_blocks(model_space, block_name: str):
    """Yield block references matching LINE_BLOCK_TEST."""
    count = com_retry(
        lambda: model_space.Count,
        "reading ModelSpace.Count",
        attempts=5,
        delay_seconds=0.4,
    )

    for i in range(count):
        try:
            entity = com_retry(
                lambda index=i: model_space.Item(index),
                f"reading ModelSpace entity {i}",
                attempts=5,
                delay_seconds=0.4,
            )
        except Exception:
            continue

        try:
            object_name = str(entity.ObjectName)
        except Exception:
            continue

        if object_name != "AcDbBlockReference":
            continue

        if get_block_name(entity) == block_name:
            yield entity


# ---------- normalization helpers ----------

def clean_value(value: Any) -> str:
    """Convert a cell or attribute value to clean string."""
    if value is None:
        return ""

    return str(value).strip()


def normalize_line_no(value: Any) -> str:
    """
    Normalize line number for matching.

    Example:
    - '4B-150-CS1'
    - '4B 150 CS1'
    - '4B_150_CS1'

    all normalize toward: '4B150CS1'
    """
    text = clean_value(value).upper()
    return re.sub(r"[^A-Z0-9]", "", text)


def normalize_compare_value(value: Any) -> str:
    """Normalize normal field values for comparison."""
    text = clean_value(value).upper()
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_row(row: Dict[str, Any]) -> Dict[str, str]:
    """Keep only expected columns and clean values."""
    return {col: clean_value(row.get(col)) for col in LINE_COLUMNS}


# ---------- P&ID extraction ----------

def extract_pid_lines(pid_path: Path) -> List[Dict[str, str]]:
    """Extract LINE_BLOCK_TEST rows from P&ID DWG."""
    if not pid_path.exists():
        print(f"ERROR: P&ID file not found: {pid_path}")
        sys.exit(1)

    acad = get_acad()

    print(f"Connected to: {acad.Caption}")
    print(f"Opening P&ID: {pid_path}")

    doc = open_drawing_with_retry(acad, pid_path)

    rows: List[Dict[str, str]] = []

    try:
        doc_name = com_retry(
            lambda: doc.Name,
            "reading active P&ID drawing name",
            attempts=5,
            delay_seconds=0.5,
        )
        print(f"Active P&ID drawing: {doc_name}")

        model_space = com_retry(
            lambda: doc.ModelSpace,
            "reading P&ID ModelSpace",
            attempts=5,
            delay_seconds=0.5,
        )

        blocks = list(find_line_blocks(model_space, LINE_BLOCK_NAME))
        print(f"Found {len(blocks)} '{LINE_BLOCK_NAME}' block(s) in P&ID.")

        for index, block_ref in enumerate(blocks, start=1):
            attrs = extract_attributes(block_ref)
            row = normalize_row(attrs)
            row["_SOURCE"] = "P&ID"
            row["_SOURCE_INDEX"] = str(index)
            rows.append(row)

    finally:
        try:
            doc.Close(False)
        except Exception:
            pass

    return rows


# ---------- Excel reading ----------

def find_latest_line_list_excel() -> Path:
    """Find newest generated clean line list Excel file."""
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

    def is_clean_line_list(path: Path) -> bool:
        name = path.name.lower()

        if name.startswith("consistency_report"):
            return False

        if "mismatch_test" in name:
            return False

        if "ai_explanation" in name:
            return False

        if "manager_summary" in name:
            return False

        return path.suffix.lower() == ".xlsx"

    candidates = [
        path for path in OUTPUT_FOLDER.glob("line_list_pid_001_*.xlsx")
        if is_clean_line_list(path)
    ]

    if not candidates:
        candidates = [
            path for path in OUTPUT_FOLDER.glob("line_list_*.xlsx")
            if is_clean_line_list(path)
        ]

    if not candidates:
        print("ERROR: No clean generated line list Excel file found in outputs folder.")
        print("Run this first:")
        print("    python -m src.use_cases.line_list_extract")
        sys.exit(1)

    latest = max(candidates, key=lambda p: p.stat().st_mtime)
    print(f"Auto-selected latest clean line list: {latest}")
    return latest


def find_header_row(ws) -> Tuple[int, Dict[str, int]]:
    """
    Find the row containing LINE_NO, SIZE, SPEC, SERVICE, FROM, TO.

    Phase 3 writes headers on row 4, but this scans to be safer.
    """
    max_scan_row = min(ws.max_row, 30)

    for row_num in range(1, max_scan_row + 1):
        values = {}

        for col_num in range(1, ws.max_column + 1):
            value = ws.cell(row=row_num, column=col_num).value
            text = clean_value(value).upper()
            if text:
                values[text] = col_num

        if all(col in values for col in LINE_COLUMNS):
            return row_num, {col: values[col] for col in LINE_COLUMNS}

    raise ValueError(
        "Could not find Excel header row with columns: "
        + ", ".join(LINE_COLUMNS)
    )


def read_excel_line_list(excel_path: Path) -> List[Dict[str, str]]:
    """Read line list rows from Excel."""
    if not excel_path.exists():
        print(f"ERROR: Excel file not found: {excel_path}")
        sys.exit(1)

    print(f"Reading Excel line list: {excel_path}")

    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active

    header_row, col_map = find_header_row(ws)

    rows: List[Dict[str, str]] = []

    for row_num in range(header_row + 1, ws.max_row + 1):
        row: Dict[str, str] = {}

        for col_name in LINE_COLUMNS:
            cell_value = ws.cell(row=row_num, column=col_map[col_name]).value
            row[col_name] = clean_value(cell_value)

        if all(not row[col] for col in LINE_COLUMNS):
            continue

        row["_SOURCE"] = "EXCEL"
        row["_SOURCE_INDEX"] = str(row_num)
        rows.append(row)

    print(f"Found {len(rows)} line row(s) in Excel.")
    return rows


# ---------- comparison ----------

def index_rows_by_line_no(
    rows: List[Dict[str, str]],
    source_name: str,
) -> Tuple[Dict[str, Dict[str, str]], List[Dict[str, str]]]:
    """
    Index rows by normalized LINE_NO.

    Returns:
    - index: normalized_line_no -> row
    - issues: missing/duplicate line number issues
    """
    grouped: Dict[str, List[Dict[str, str]]] = {}
    issues: List[Dict[str, str]] = []

    for row in rows:
        line_no = row.get("LINE_NO", "")
        key = normalize_line_no(line_no)

        if not key:
            issues.append(
                {
                    "type": "MISSING_LINE_NO",
                    "line_no": line_no,
                    "field": "LINE_NO",
                    "pid_value": "",
                    "excel_value": "",
                    "details": f"{source_name} row has no LINE_NO. Source index: {row.get('_SOURCE_INDEX', '')}",
                }
            )
            continue

        grouped.setdefault(key, []).append(row)

    index: Dict[str, Dict[str, str]] = {}

    for key, group in grouped.items():
        if len(group) > 1:
            issues.append(
                {
                    "type": "DUPLICATE_LINE_NO",
                    "line_no": group[0].get("LINE_NO", key),
                    "field": "LINE_NO",
                    "pid_value": "",
                    "excel_value": "",
                    "details": f"{source_name} has {len(group)} rows for normalized line key '{key}'.",
                }
            )

        index[key] = group[0]

    return index, issues


def compare_pid_vs_excel(
    pid_rows: List[Dict[str, str]],
    excel_rows: List[Dict[str, str]],
) -> List[Dict[str, str]]:
    """Compare P&ID rows against Excel rows."""
    pid_index, pid_issues = index_rows_by_line_no(pid_rows, "P&ID")
    excel_index, excel_issues = index_rows_by_line_no(excel_rows, "EXCEL")

    mismatches: List[Dict[str, str]] = []
    mismatches.extend(pid_issues)
    mismatches.extend(excel_issues)

    pid_keys = set(pid_index.keys())
    excel_keys = set(excel_index.keys())

    for key in sorted(pid_keys - excel_keys):
        pid_row = pid_index[key]
        mismatches.append(
            {
                "type": "MISSING_IN_EXCEL",
                "line_no": pid_row.get("LINE_NO", key),
                "field": "LINE_NO",
                "pid_value": pid_row.get("LINE_NO", ""),
                "excel_value": "",
                "details": "Line exists in P&ID but not in Excel line list.",
            }
        )

    for key in sorted(excel_keys - pid_keys):
        excel_row = excel_index[key]
        mismatches.append(
            {
                "type": "MISSING_IN_PID",
                "line_no": excel_row.get("LINE_NO", key),
                "field": "LINE_NO",
                "pid_value": "",
                "excel_value": excel_row.get("LINE_NO", ""),
                "details": "Line exists in Excel line list but not in P&ID.",
            }
        )

    for key in sorted(pid_keys & excel_keys):
        pid_row = pid_index[key]
        excel_row = excel_index[key]

        for field in COMPARE_COLUMNS:
            pid_value = pid_row.get(field, "")
            excel_value = excel_row.get(field, "")

            if normalize_compare_value(pid_value) != normalize_compare_value(excel_value):
                mismatches.append(
                    {
                        "type": "FIELD_MISMATCH",
                        "line_no": pid_row.get("LINE_NO") or excel_row.get("LINE_NO") or key,
                        "field": field,
                        "pid_value": pid_value,
                        "excel_value": excel_value,
                        "details": f"{field} differs between P&ID and Excel.",
                    }
                )

    return mismatches


# ---------- report output ----------

def style_header_row(ws, row_num: int):
    """Apply simple header styling."""
    fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="center", vertical="center")

    for cell in ws[row_num]:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


def auto_width(ws):
    """Auto-fit column widths."""
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 10

        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            value_len = len(str(value or ""))

            if value_len > max_len:
                max_len = value_len

        ws.column_dimensions[letter].width = min(max_len + 4, 60)


def write_rows_sheet(wb, title: str, rows: List[Dict[str, str]]):
    """Write extracted rows into a worksheet."""
    ws = wb.create_sheet(title)

    headers = LINE_COLUMNS + ["_SOURCE_INDEX"]
    ws.append(headers)
    style_header_row(ws, 1)

    for row in rows:
        ws.append([row.get(col, "") for col in headers])

    auto_width(ws)
    ws.freeze_panes = "A2"


def write_report(
    report_path: Path,
    pid_path: Path,
    excel_path: Path,
    pid_rows: List[Dict[str, str]],
    excel_rows: List[Dict[str, str]],
    mismatches: List[Dict[str, str]],
):
    """Write mismatch report to Excel."""
    report_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"

    result = "PASS" if not mismatches else "FAIL"

    summary_rows = [
        ("Consistency Check Result", result),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("P&ID DWG", str(pid_path)),
        ("Excel Line List", str(excel_path)),
        ("P&ID Rows", len(pid_rows)),
        ("Excel Rows", len(excel_rows)),
        ("Mismatch Count", len(mismatches)),
    ]

    for row in summary_rows:
        ws.append(row)

    ws["A1"].font = Font(bold=True, size=14)
    ws["B1"].font = Font(bold=True, size=14)
    ws["B1"].fill = PatternFill(
        start_color="C6EFCE" if result == "PASS" else "FFC7CE",
        end_color="C6EFCE" if result == "PASS" else "FFC7CE",
        fill_type="solid",
    )

    auto_width(ws)

    # Mismatch sheet
    ws_m = wb.create_sheet("Mismatches")
    mismatch_headers = ["type", "line_no", "field", "pid_value", "excel_value", "details"]
    ws_m.append(mismatch_headers)
    style_header_row(ws_m, 1)

    if mismatches:
        for item in mismatches:
            ws_m.append([item.get(header, "") for header in mismatch_headers])
    else:
        ws_m.append(["NO_MISMATCHES", "", "", "", "", "P&ID and Excel line list match."])

    auto_width(ws_m)
    ws_m.freeze_panes = "A2"

    write_rows_sheet(wb, "P&ID Lines", pid_rows)
    write_rows_sheet(wb, "Excel Lines", excel_rows)

    wb.save(report_path)


# ---------- main ----------

@log_job("consistency_check")
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--pid",
        type=str,
        default=str(DEFAULT_PID_DWG),
        help="Path to P&ID DWG file.",
    )
    parser.add_argument(
        "--excel",
        type=str,
        default="",
        help="Path to line list Excel file. If omitted, latest clean generated line list is used.",
    )
    parser.add_argument(
        "--report",
        type=str,
        default="",
        help="Output report path. If omitted, timestamped report is written to outputs/.",
    )

    args = parser.parse_args()

    pid_path = Path(args.pid)

    if args.excel.strip():
        excel_path = Path(args.excel)
    else:
        excel_path = find_latest_line_list_excel()

    if args.report.strip():
        report_path = Path(args.report)
    else:
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        report_path = OUTPUT_FOLDER / f"consistency_report_pid_001_{timestamp}.xlsx"

    print("=" * 80)
    print("PHASE 6 CONSISTENCY CHECK — P&ID vs Excel Line List")
    print("=" * 80)

    pid_rows = extract_pid_lines(pid_path)
    excel_rows = read_excel_line_list(excel_path)

    mismatches = compare_pid_vs_excel(pid_rows, excel_rows)

    print("\nComparison summary:")
    print(f"  P&ID rows:       {len(pid_rows)}")
    print(f"  Excel rows:      {len(excel_rows)}")
    print(f"  Mismatches:      {len(mismatches)}")

    if mismatches:
        print("\nMismatches found:")
        for item in mismatches:
            print(
                f"  - {item['type']} | line={item['line_no']} | "
                f"field={item['field']} | P&ID={item['pid_value']!r} | "
                f"Excel={item['excel_value']!r}"
            )
    else:
        print("\nNo mismatches found. P&ID and Excel line list are consistent.")

    write_report(
        report_path=report_path,
        pid_path=pid_path,
        excel_path=excel_path,
        pid_rows=pid_rows,
        excel_rows=excel_rows,
        mismatches=mismatches,
    )

    print(f"\nReport written to: {report_path}")

    if mismatches:
        print("\nResult: FAIL — review the mismatch report.")
    else:
        print("\nResult: PASS — baseline consistency check is working.")


if __name__ == "__main__":
    main()
