"""Quick test that openpyxl is installed and can write a real Excel file."""

from pathlib import Path
from openpyxl import Workbook


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Hello"

    ws["A1"] = "Hello from Python"
    ws["B1"] = 42
    ws.append(["row 2 col A", "row 2 col B"])
    ws.append(["row 3 col A", 3.14])

    out_path = Path(__file__).parent.parent.parent / "excel_hello.xlsx"
    wb.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()