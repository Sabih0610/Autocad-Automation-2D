"""
Create a deliberately incorrect copy of the latest generated line list.

This is for Phase 6 mismatch testing.

It does NOT touch AutoCAD.
It does NOT modify the original Excel file.

Run:
    python -m src.scratch.make_mismatch_line_list

Then copy the printed command and run it.
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from src.use_cases.consistency_check import (
    OUTPUT_FOLDER,
    find_header_row,
    find_latest_line_list_excel,
)


def main():
    source_path = find_latest_line_list_excel()

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    mismatch_path = OUTPUT_FOLDER / f"line_list_pid_001_MISMATCH_TEST_{timestamp}.xlsx"

    shutil.copy2(source_path, mismatch_path)

    wb = load_workbook(mismatch_path)
    ws = wb.active

    header_row, col_map = find_header_row(ws)

    first_data_row = header_row + 1
    second_data_row = header_row + 2

    # Deliberate mismatches
    ws.cell(row=first_data_row, column=col_map["SERVICE"]).value = "WRONG SERVICE"
    ws.cell(row=second_data_row, column=col_map["SIZE"]).value = "999 inch"

    wb.save(mismatch_path)

    print("Mismatch test file created:")
    print(mismatch_path)

    print("\nNow run this command:")
    print(f'python -m src.use_cases.consistency_check --excel "{mismatch_path}"')


if __name__ == "__main__":
    main()